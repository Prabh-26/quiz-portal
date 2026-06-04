from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "Online_Quiz_Portal_Dissertation.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(header_cells[i], header, bold=True)
        set_cell_shading(header_cells[i], "E8EEF5")
        if widths:
            header_cells[i].width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], str(value))
            if widths:
                cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(46, 116, 181) if level <= 2 else RGBColor(31, 77, 120)
    return p


def add_para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(4)


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)

for style_name, size, color in [
    ("Title", 20, RGBColor(11, 37, 69)),
    ("Heading 1", 16, RGBColor(46, 116, 181)),
    ("Heading 2", 13, RGBColor(46, 116, 181)),
    ("Heading 3", 12, RGBColor(31, 77, 120)),
]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = color

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A Dissertation Report\n")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(11, 37, 69)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ON\n")
r.bold = True
r.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ONLINE QUIZ PORTAL")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(31, 60, 136)

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Submitted in partial fulfillment of the requirements for the final year project/dissertation").italic = True

doc.add_paragraph()
meta = doc.add_table(rows=6, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
cover_rows = [
    ("Submitted By", "[Your Name]"),
    ("Roll Number", "[Your Roll Number]"),
    ("Course", "[Your Course]"),
    ("Submitted To", "[Guide/Faculty Name]"),
    ("Department", "[Department Name]"),
    ("Institution", "[College/School Name]"),
]
for i, row in enumerate(cover_rows):
    set_cell_text(meta.rows[i].cells[0], row[0], bold=True)
    set_cell_text(meta.rows[i].cells[1], row[1])
    set_cell_shading(meta.rows[i].cells[0], "E8EEF5")

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Academic Year: 2025-2026").bold = True

page_break(doc)

# Certificate
add_heading(doc, "Certificate", 1)
add_para(
    doc,
    "This is to certify that the dissertation project titled \"Online Quiz Portal\" has been carried out by "
    "[Your Name], bearing roll number [Your Roll Number], under the guidance of [Guide/Faculty Name]. "
    "The work submitted in this report is a record of the project developed for academic purpose.",
)
doc.add_paragraph("\n\n")
add_table(
    doc,
    ["Name", "Signature", "Date"],
    [
        ["Project Guide", "", ""],
        ["Head of Department", "", ""],
        ["External Examiner", "", ""],
    ],
    [2.2, 2.2, 1.4],
)

add_heading(doc, "Acknowledgement", 1)
add_para(
    doc,
    "I would like to express my sincere gratitude to my project guide, faculty members, and institution "
    "for their guidance and support throughout the development of this project. I am also thankful to my "
    "friends and classmates for their suggestions during testing and improvement of the application.",
)
page_break(doc)

# Abstract and TOC-like outline
add_heading(doc, "Abstract", 1)
add_para(
    doc,
    "The Online Quiz Portal is a web-based application developed using Java Spring Boot, MySQL, HTML, CSS, "
    "and JavaScript. The project provides a digital platform where administrators can manage quiz questions "
    "and tests, while students can register, log in, attempt quizzes, and view their previous results. "
    "The system includes role-based access, approved roll number registration, course-wise test display, "
    "score calculation, and result history. The application reduces manual effort in conducting quizzes and "
    "provides a structured way to track student performance."
)

add_heading(doc, "Contents", 1)
contents = [
    "1. Introduction",
    "2. Objectives",
    "3. Existing System and Proposed System",
    "4. Technology Stack",
    "5. System Analysis",
    "6. Database Design",
    "7. Module Description",
    "8. Implementation Details",
    "9. Testing",
    "10. Limitations and Future Scope",
    "11. Conclusion",
    "12. References",
]
add_numbered(doc, contents)
page_break(doc)

# Main report
add_heading(doc, "1. Introduction", 1)
add_para(
    doc,
    "Educational institutions frequently conduct tests to evaluate student understanding. Traditional "
    "paper-based quizzes require manual question preparation, answer checking, result calculation, and "
    "record maintenance. The Online Quiz Portal solves these problems by providing a web-based system for "
    "quiz management and automated score calculation."
)
add_para(
    doc,
    "The project supports two main user roles: administrator and student. Administrators manage questions "
    "and tests. Students register through approved roll numbers, log in, view course-wise available tests, "
    "attempt quizzes, and check their result history."
)

add_heading(doc, "2. Objectives", 1)
add_bullets(
    doc,
    [
        "To develop a web-based quiz system using Java Spring Boot and MySQL.",
        "To provide separate login flows for administrator and student users.",
        "To allow only approved students to register using valid roll numbers.",
        "To prevent one roll number from being used by multiple accounts.",
        "To allow administrators to add, view, and delete quiz questions.",
        "To display course-wise tests on the student dashboard.",
        "To calculate scores automatically and store results in the database.",
        "To provide students with previous test result history.",
    ],
)

add_heading(doc, "3. Existing System and Proposed System", 1)
add_heading(doc, "3.1 Existing System", 2)
add_para(
    doc,
    "In the traditional system, quizzes are conducted manually. Teachers prepare question papers, students "
    "answer on paper, and evaluation is done manually. This process is time-consuming and prone to calculation "
    "or record-keeping errors."
)
add_heading(doc, "3.2 Proposed System", 2)
add_para(
    doc,
    "The proposed Online Quiz Portal automates quiz creation, student registration validation, test display, "
    "score calculation, and result storage. It provides a faster, cleaner, and more organized way to conduct "
    "academic quizzes."
)

add_heading(doc, "4. Technology Stack", 1)
add_table(
    doc,
    ["Layer", "Technology", "Purpose"],
    [
        ["Frontend", "HTML, CSS, JavaScript", "User interface and browser-side interaction"],
        ["Backend", "Java Spring Boot", "Handles requests, business logic, APIs"],
        ["Database", "MySQL", "Stores users, questions, tests, courses, and results"],
        ["ORM", "Spring Data JPA", "Maps Java classes to database tables"],
        ["IDE", "IntelliJ IDEA Community", "Code development and project management"],
    ],
    [1.5, 2.0, 3.0],
)

add_heading(doc, "5. System Analysis", 1)
add_heading(doc, "5.1 User Roles", 2)
add_bullets(
    doc,
    [
        "Admin: Logs in using an admin account and manages quiz questions and tests.",
        "Student: Registers using an approved roll number, logs in, attempts tests, and views results.",
    ],
)
add_heading(doc, "5.2 Authentication and Authorization", 2)
add_para(
    doc,
    "Authentication verifies the user through username and password. Authorization is handled using the role "
    "stored in the users table. The role decides whether the user should access the admin panel or student "
    "dashboard."
)

add_heading(doc, "6. Database Design", 1)
add_para(doc, "The database used for the project is named quiz_portal. The main tables are listed below.")
add_table(
    doc,
    ["Table", "Purpose"],
    [
        ["users", "Stores registered users, roles, roll numbers, and course IDs"],
        ["approved_students", "Stores institution-approved roll numbers for student registration"],
        ["courses", "Stores course names such as BCA, B.Tech CSE, MCA"],
        ["subjects", "Stores subjects mapped to courses"],
        ["tests", "Stores test titles mapped to subjects and courses"],
        ["questions", "Stores MCQ questions, options, correct answer, and test ID"],
        ["results", "Stores student score, total marks, test attempt time, and user reference"],
    ],
    [2.1, 4.4],
)

add_heading(doc, "6.1 Important Database Relationships", 2)
add_bullets(
    doc,
    [
        "approved_students.course_id identifies the course of an approved student.",
        "users.roll_number is unique, so one roll number cannot register twice.",
        "users.course_id stores the registered student's course.",
        "subjects.course_id connects subjects to a course.",
        "tests.subject_id and tests.course_id connect a test to a subject and course.",
        "questions.test_id connects questions to a specific test.",
        "results.user_id connects quiz results to the student account.",
    ],
)

add_heading(doc, "7. Module Description", 1)
add_heading(doc, "7.1 Admin Module", 2)
add_bullets(
    doc,
    [
        "Admin login using username and password.",
        "Admin panel access protection using role validation.",
        "Add new MCQ questions with four options and correct answer.",
        "View all available questions.",
        "Delete questions from the system.",
        "View total number of questions.",
    ],
)

add_heading(doc, "7.2 Student Registration Module", 2)
add_bullets(
    doc,
    [
        "Student enters username, password, and roll number.",
        "System checks if roll number exists in approved_students.",
        "System checks if roll number is already used in users.",
        "If valid, account is created with role student.",
        "Student course_id is copied from approved_students to users.",
    ],
)

add_heading(doc, "7.3 Student Dashboard Module", 2)
add_bullets(
    doc,
    [
        "Student lands on dashboard after login.",
        "Dashboard fetches tests based on student's course_id.",
        "Available tests are displayed to the student.",
        "Student can navigate to quiz page and result history.",
    ],
)

add_heading(doc, "7.4 Quiz and Result Module", 2)
add_bullets(
    doc,
    [
        "Questions are loaded from the backend.",
        "Student selects answers using radio buttons.",
        "System validates that all questions are answered.",
        "Score is calculated automatically.",
        "Result is saved in MySQL and shown in result history.",
    ],
)

add_heading(doc, "8. Implementation Details", 1)
add_heading(doc, "8.1 Project Structure", 2)
add_table(
    doc,
    ["Package/Folder", "Description"],
    [
        ["model", "Contains entity classes such as User, Question, Result, Course, Subject, and Test"],
        ["repository", "Contains Spring Data JPA interfaces for database operations"],
        ["service", "Contains business logic such as login, registration, and fetching tests"],
        ["controller", "Contains REST API endpoints"],
        ["static", "Contains HTML, CSS, and JavaScript pages"],
    ],
    [2.0, 4.5],
)

add_heading(doc, "8.2 Main API Endpoints", 2)
add_table(
    doc,
    ["Endpoint", "Method", "Purpose"],
    [
        ["/login", "POST", "Verifies username/password and returns role"],
        ["/register", "POST", "Registers student after approved roll number validation"],
        ["/questions", "GET", "Fetches all questions or questions by testId"],
        ["/questions/add", "POST", "Adds a new question"],
        ["/questions/delete/{id}", "DELETE", "Deletes a question"],
        ["/result/save", "POST", "Saves quiz result"],
        ["/result/history", "GET", "Shows previous results of student"],
        ["/subjects/my", "GET", "Fetches subjects for logged-in student's course"],
        ["/tests/my", "GET", "Fetches tests for logged-in student's course"],
    ],
    [2.1, 1.0, 3.4],
)

add_heading(doc, "8.3 Application Flow", 2)
add_numbered(
    doc,
    [
        "Admin account is created by the institution or IT department.",
        "Approved student roll numbers are inserted into the approved_students table.",
        "Student registers using username, password, and roll number.",
        "System validates the roll number and stores the student course.",
        "Student logs in and reaches the dashboard.",
        "Dashboard shows tests available for the student's course.",
        "Student attempts quiz and submits answers.",
        "System calculates score and stores the result.",
        "Student can view result history from the result page.",
    ],
)

add_heading(doc, "9. Testing", 1)
add_table(
    doc,
    ["Test Case", "Input/Action", "Expected Result", "Status"],
    [
        ["Admin login", "admin/admin123", "Admin panel opens", "Passed"],
        ["Student registration", "Valid roll number 103", "Registration successful", "Passed"],
        ["Duplicate roll number", "Use same roll number again", "Registration failed", "Passed"],
        ["Invalid roll number", "Use 999", "Registration failed", "Passed"],
        ["Add question", "Enter question and options", "Question saved", "Passed"],
        ["Delete question", "Click delete", "Question removed", "Passed"],
        ["Student dashboard", "Login as student", "Course-wise tests displayed", "Passed"],
        ["Quiz validation", "Submit unanswered quiz", "Warning displayed", "Passed"],
        ["Result save", "Submit quiz", "Result stored in MySQL", "Passed"],
    ],
    [1.7, 1.8, 2.1, 0.9],
)

add_heading(doc, "10. Limitations and Future Scope", 1)
add_heading(doc, "10.1 Limitations", 2)
add_bullets(
    doc,
    [
        "Current password storage is plain text and should be improved using password hashing.",
        "Frontend localStorage is used for simple role/session handling; backend session security can be added.",
        "AI question generation is planned but not yet fully integrated.",
        "The current UI is functional and basic; it can be further improved.",
    ],
)
add_heading(doc, "10.2 Future Scope", 2)
add_bullets(
    doc,
    [
        "Add AI-based MCQ generation for admins using Gemini or another free-tier AI API.",
        "Add admin interface for creating subjects and tests directly from the browser.",
        "Add detailed student performance charts and analytics.",
        "Add Spring Security for professional authentication and authorization.",
        "Add timer-based tests and automatic submission.",
        "Add email or OTP verification for student registration.",
    ],
)

add_heading(doc, "11. Conclusion", 1)
add_para(
    doc,
    "The Online Quiz Portal successfully provides a digital platform for conducting quizzes in an educational "
    "environment. It supports admin and student roles, approved roll number registration, course-wise test "
    "display, automatic score calculation, and result history. The project demonstrates practical use of Java "
    "Spring Boot, MySQL, REST APIs, and frontend technologies. The system can be expanded further with AI-based "
    "question generation, advanced security, and performance analytics."
)

add_heading(doc, "12. References", 1)
add_bullets(
    doc,
    [
        "Spring Boot Documentation",
        "Spring Data JPA Documentation",
        "MySQL Documentation",
        "Java Documentation",
        "HTML, CSS, and JavaScript Web References",
    ],
)

doc.core_properties.title = "Online Quiz Portal Dissertation"
doc.core_properties.subject = "Final Year Project Report"
doc.core_properties.author = "[Your Name]"
doc.save(OUT)
print(OUT)
